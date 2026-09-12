#!/usr/bin/env python3

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "chapter-01-gem5-systemc-foundation.md"
OUTPUT = ROOT / "docs" / "zhihu" / "chapter-01-gem5-systemc-foundation.docx"
BODY_FONT = "Arial Unicode MS"


def set_run_font(run, latin=BODY_FONT, east_asia=BODY_FONT, size=None):
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), latin)
    r_pr.rFonts.set(qn("w:hAnsi"), latin)
    r_pr.rFonts.set(qn("w:eastAsia"), east_asia)
    r_pr.rFonts.set(qn("w:hint"), "eastAsia")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_inline(paragraph, text, size=10.5, color="20242A"):
    token = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")
    position = 0
    for match in token.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size)
            run.font.color.rgb = RGBColor.from_string(color)
        value = match.group(0)
        if value.startswith("**"):
            run = paragraph.add_run(value[2:-2])
            set_run_font(run, size=size)
            run.bold = True
        elif value.startswith("`"):
            run = paragraph.add_run(value[1:-1])
            set_run_font(run, latin="Menlo", size=size - 0.5)
            run.font.color.rgb = RGBColor(32, 78, 121)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", value).groups()
            display = label if label == url else f"{label}（{url}）"
            run = paragraph.add_run(display)
            set_run_font(run, size=size)
            run.font.color.rgb = RGBColor(5, 99, 193)
            run.underline = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size)
        run.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 36, 42)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(7)

    title = doc.styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(21)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 18, 8),
        ("Heading 2", 13, 13, 5),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_code_block(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F5F7")
    p_pr.append(shd)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, latin="Menlo", size=8.7)
    run.font.color.rgb = RGBColor(35, 45, 55)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.autofit = True
    set_table_borders(table)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, size=9.5)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cell, "F3F7FA")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_markdown(doc, text):
    lines = text.splitlines()
    index = 0
    code = None
    table_rows = []
    first_heading = True
    in_references = False

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if code is None:
                code = []
            else:
                add_code_block(doc, code)
                code = None
            index += 1
            continue
        if code is not None:
            code.append(line)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                index += 1
                continue
            table_rows.append(cells)
            next_line = lines[index + 1].strip() if index + 1 < len(lines) else ""
            if not (next_line.startswith("|") and next_line.endswith("|")):
                add_table(doc, table_rows)
                table_rows = []
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = SOURCE.parent / image_match.group(2)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = True
            if "04-lifecycle" in image_path.name:
                width = Inches(4.9)
            elif image_path.name.startswith("02-"):
                width = Inches(5.2)
            elif image_path.name.startswith("03-"):
                width = Inches(5.4)
            else:
                width = Inches(6.35)
            paragraph.add_run().add_picture(str(image_path), width=width)
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            value = heading.group(2)
            if first_heading and level == 1:
                paragraph = doc.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, value, size=21, color="000000")
                first_heading = False
            else:
                paragraph = doc.add_heading(level=level)
                add_inline(paragraph, value, size={1: 17, 2: 13, 3: 11.5}[level], color="000000")
            in_references = value == "参考资料"
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(14)
            run = paragraph.add_run(stripped[2:])
            set_run_font(run, size=11)
            run.italic = True
            run.font.color.rgb = RGBColor(89, 98, 108)
            index += 1
            continue

        list_match = re.match(r"^([-*]|\d+\.)\s+(.+)$", stripped)
        if list_match:
            marker = list_match.group(1)
            if in_references and marker.endswith("."):
                paragraph = doc.add_paragraph()
                add_inline(paragraph, f"{marker} {list_match.group(2)}")
                paragraph.paragraph_format.left_indent = Cm(0.45)
                paragraph.paragraph_format.first_line_indent = Cm(-0.45)
                paragraph.paragraph_format.space_after = Pt(4)
                index += 1
                continue
            style = "List Bullet" if marker in ("-", "*") else "List Number"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, list_match.group(2))
            index += 1
            continue

        if re.match(r"^图\s*\d+", stripped):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(10)
            run = paragraph.add_run(stripped)
            set_run_font(run, size=9)
            run.font.color.rgb = RGBColor(89, 98, 108)
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1


def main():
    document = Document()
    configure_document(document)
    parse_markdown(document, SOURCE.read_text(encoding="utf-8"))

    core = document.core_properties
    core.title = "从零搭建 gem5 与 SystemC 联合仿真环境"
    core.subject = "NPU 系统建模与联合仿真系列第一章"
    core.author = "Ch'in"
    core.keywords = "NPU, gem5, SystemC, RISC-V, cycle model"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
